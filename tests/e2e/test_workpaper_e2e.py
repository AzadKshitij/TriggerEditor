"""Workpaper end to end: Audit node + File Outputs + Workpaper Output, then the export.

A real window runs a real workflow on a worker thread:
ledger.csv -> Duplicates -> (D) Workpaper Output exceptions, (S) summary,
ledger.csv -> Workpaper Output population, (U) File Output as Excel.
"""

from __future__ import annotations

import hashlib
from pathlib import Path

import openpyxl
import polars as pl
import pytest
from nodeeditor.node_edge import Edge

from tests.support.audit_helpers import newest_record, store_dir_of
from tests.support.audit_ledger import make_ledger
from tests.support.workflows import run_workflow_and_wait
from tests.support.workpaper import load
from trigger_designer.core.audit.store import verify_store
from trigger_designer.core.node_configuration import (
    AuditNodes,
    IONodes,
    NodeTypes,
    ReportNodes,
    get_class_from_opcode,
)
from trigger_designer.qt.design_window import TriggerSubWindow
from trigger_designer.qt.helpers import workpaper_export

pytestmark = pytest.mark.e2e

KEYS = ["vendor", "invoice_no", "amount"]


@pytest.fixture(autouse=True)
def _no_dialogs(no_modal_dialogs):
    return no_modal_dialogs


@pytest.fixture(scope="module")
def ledger_csv(tmp_path_factory):
    path = tmp_path_factory.mktemp("ledger") / "ledger.csv"
    make_ledger().frame.write_csv(path)
    return path


def _build(root: Path, source: Path, **wp_settings):
    out = root / "out"
    out.mkdir()
    window = TriggerSubWindow()
    scene = window.scene
    file_input = get_class_from_opcode(IONodes.FILE_INPUT, NodeTypes.IO)(scene)
    file_input.content.filePath = str(source)
    file_input.content.file_type = "csv"
    dups = get_class_from_opcode(AuditNodes.DUPLICATES, NodeTypes.AUDIT)(scene)
    dups.content.key_columns = list(KEYS)
    csv_out = get_class_from_opcode(IONodes.FILE_OUTPUT, NodeTypes.IO)(scene)
    csv_out.content.settings.update(filePath=str(out / "dups.csv"), file_format="csv")
    xlsx_out = get_class_from_opcode(IONodes.FILE_OUTPUT, NodeTypes.IO)(scene)
    xlsx_out.content.settings.update(
        filePath=str(out / "unique.xlsx"), file_format="excel"
    )
    wp = get_class_from_opcode(IONodes.WORKPAPER_OUTPUT, NodeTypes.IO)(scene)
    wp.content.settings.update(
        output_path=str(out / "wp.xlsx"),
        client="Acme Ltd",
        engagement="FY2025 audit",
        period="Year ended 31 December 2025",
        preparer="A. Auditor",
        prepared_date="2026-03-31",
        objective="Show that no invoice was paid twice.",
        procedure="Matched vendor, invoice number and amount.",
        amount_column="amount",
        **wp_settings,
    )
    note = get_class_from_opcode(ReportNodes.COMMENT, NodeTypes.REPORT)(scene)
    note.content.settings["text"] = "Duplicates reviewed with the AP manager."
    Edge(scene, file_input.outputs[0], dups.inputs[0])
    Edge(scene, dups.outputs[0], csv_out.inputs[0])
    Edge(scene, dups.outputs[1], xlsx_out.inputs[0])
    Edge(scene, dups.outputs[0], wp.inputs[0])
    Edge(scene, file_input.outputs[0], wp.inputs[1])
    Edge(scene, dups.outputs[2], wp.inputs[2])
    tds = root / "wp_flow.tds"
    scene.saveToFile(str(tds))
    window.close()
    return tds, out


def _run(qtbot, root, source, **wp_settings):
    tds, out = _build(root, source, **wp_settings)
    window = TriggerSubWindow()
    qtbot.addWidget(window)
    assert window.fileLoad(str(tds)) is True
    assert window.validateLoadedWorkflow() == []
    run_workflow_and_wait(qtbot, window)
    assert not any(node.isInvalid() for node in window.getAllNodes())
    return window, tds, out


def _population_rows(record):
    for node in record["nodes"]:
        summary = node.get("test_summary")
        if summary:
            return summary["population_rows"]
    return None


def test_workpaper_from_a_real_run(qtbot, tmp_path, ledger_csv, monkeypatch):
    window, tds, out = _run(qtbot, tmp_path, ledger_csv)
    dups = pl.read_csv(out / "dups.csv", infer_schema=False)
    ledger_rows = pl.read_csv(ledger_csv, infer_schema=False).height
    n = dups.height
    assert n > 0

    wb = load(out / "wp.xlsx")
    assert wb.sheetnames == ["Summary", "Exceptions", "Run log", "Workflow"]
    summary = wb["Summary"]
    cached = load(out / "wp.xlsx", data_only=True)["Summary"]
    assert summary["B6"].value == "Acme Ltd" and summary["B10"].value == "A. Auditor"
    assert summary["B11"].value.date().isoformat() == "2026-03-31"
    assert cached["B22"].value == ledger_rows
    assert cached["B24"].value == n and cached["B23"].value == n
    assert summary["B26"].value == f"Agrees to the test node ({n:,} exceptions)"
    assert summary["F22"].value == "Counted from the Population input"
    assert summary["B9"].value  # the test name came from the Summary input
    assert summary["B25"].value.startswith("Not needed")
    exceptions = wb["Exceptions"]
    assert exceptions.max_row == n + 2
    assert [c.value for c in exceptions[2]][:3] == [
        "Ref",
        "Tick mark",
        "Auditor comment",
    ]
    assert exceptions.freeze_panes == "B3"

    column_a = [c.value for row in summary.iter_rows(min_col=1, max_col=1) for c in row]
    assert ledger_csv.as_posix() in column_a  # record paths use forward slashes
    ledger_sha = hashlib.sha256(ledger_csv.read_bytes()).hexdigest()
    column_f = [c.value for row in summary.iter_rows(min_col=6, max_col=6) for c in row]
    assert ledger_sha in column_f

    run_log = wb["Run log"]
    assert run_log["A2"].value.startswith(
        "Snapshot taken while the workflow was running"
    )
    steps = [c.value for c in run_log["B"] if c.value]
    assert "Duplicates" in steps and "File Input" in steps

    flow = wb["Workflow"]
    titles = [c.value for c in flow["B"]]
    assert "Duplicates" in titles and "File Input" in titles
    assert "Duplicates reviewed with the AP manager." in [
        c.value for row in flow.iter_rows(min_col=2, max_col=2) for c in row
    ]

    # File Output as Excel streams the unique rows into a readable workbook
    unique = openpyxl.load_workbook(out / "unique.xlsx").active
    assert unique.max_row == ledger_rows - n + 1
    assert unique.freeze_panes == "A2" and unique.auto_filter.ref

    # the run record is sealed, verifies, and fingerprints the workpaper
    directory = store_dir_of(tds)
    assert verify_store(directory).ok
    record = newest_record(directory)
    assert record["status"] == "ok"
    ops = [node["op"] for node in record["nodes"]]
    assert "input.workpaper_output" in ops and "audit.duplicates" in ops
    fingerprints = {Path(o["path"]).name: o for o in record["outputs"]}
    workbook_sha = hashlib.sha256((out / "wp.xlsx").read_bytes()).hexdigest()
    assert fingerprints["wp.xlsx"]["sha256"] == workbook_sha
    assert "unique.xlsx" in fingerprints

    # File > Export Workpaper... builds the whole-run workpaper from that record
    class Main:
        def getCurrentNodeEditorWidget(self):
            return window

    class Dialog:
        def __init__(self, _parent, _suggested):
            pass

        def exec(self):
            return workpaper_export.QDialog.DialogCode.Accepted

        def settings(self):
            return {"output_path": str(out / "run.xlsx"), "client": "Acme Ltd"}

    shown = []
    monkeypatch.setattr(
        workpaper_export.QMessageBox, "information", lambda *a, **k: shown.append(a[2])
    )
    path = workpaper_export.export_workpaper(Main(), dialog_class=Dialog)
    assert path and Path(path).name == "run.xlsx" and shown
    run_wb = load(path)
    assert run_wb.sheetnames == ["Summary", "Run log", "Workflow"]
    rows = {
        row[0].value: row
        for row in run_wb["Summary"].iter_rows(min_row=22, max_row=30, max_col=6)
        if row[0].value
    }
    test_row = rows["Duplicates"]
    assert test_row[3].value == n and test_row[2].value == _population_rows(record)
    assert run_wb["Run log"]["A2"].value == "Completed run."
    assert verify_store(directory).ok  # exporting never touches the sealed history


def test_second_run_creates_a_numbered_copy_not_an_overwrite(
    qtbot, tmp_path, ledger_csv
):
    window, _tds, out = _run(qtbot, tmp_path, ledger_csv)
    first = (out / "wp.xlsx").read_bytes()
    run_workflow_and_wait(qtbot, window)
    assert (out / "wp.xlsx").read_bytes() == first
    assert (out / "wp (2).xlsx").exists()


def test_replace_mode_overwrites_the_workpaper(qtbot, tmp_path, ledger_csv):
    window, _tds, out = _run(qtbot, tmp_path, ledger_csv, if_exists="replace")
    run_workflow_and_wait(qtbot, window)
    assert not (out / "wp (2).xlsx").exists() and (out / "wp.xlsx").exists()


def test_row_limit_truncates_and_writes_the_full_csv(qtbot, tmp_path, ledger_csv):
    _window, _tds, out = _run(qtbot, tmp_path, ledger_csv, row_limit=2)
    n = pl.read_csv(out / "dups.csv", infer_schema=False).height
    assert n > 2
    wb = load(out / "wp.xlsx")
    assert wb["Exceptions"]["A1"].value.startswith(
        f"TRUNCATED - showing the first 2 of {n}"
    )
    full = out / "wp_exceptions_full.csv"
    assert full.exists()
    assert pl.read_csv(full, infer_schema=False).height == n
