"""Shared helpers for workpaper tests: read workbooks back with openpyxl."""

from __future__ import annotations

import datetime as dt

import openpyxl
import polars as pl

from trigger_designer.core.workpaper.spec import WorkpaperSpec

FIXED_NOW = dt.datetime(2026, 3, 31, 9, 30)

EXCEPTIONS = pl.DataFrame(
    {
        "Invoice": ["INV006", "INV008", "INV010"],
        "Vendor": ["Zenith", "Acme", "Borealis"],
        "Amount": [1500.75, 3000.0, 2200.0],
        "Note": ["Above limit", "=1+1", "Above limit"],
    }
)
POPULATION = pl.DataFrame(
    {
        "Invoice": [f"INV{i:03d}" for i in range(1, 11)],
        "Amount": [
            100.0,
            250.5,
            1200.0,
            75.25,
            980.0,
            1500.75,
            40.0,
            3000.0,
            610.1,
            2200.0,
        ],
    }
)


def load(path, data_only=False):
    return openpyxl.load_workbook(path, data_only=data_only)


def find_label(ws, label, col=1):
    for r in range(1, ws.max_row + 1):
        if ws.cell(r, col).value == label:
            return r
    raise AssertionError(f"label {label!r} not found in {ws.title}")


def value_right_of(ws, label):
    return ws.cell(find_label(ws, label), 2).value


def make_spec(tmp_path, **kw):
    base = {
        "output_path": str(tmp_path / "wp.xlsx"),
        "title": "Invoice approval test",
        "client": "Acme Ltd",
        "engagement": "FY2025 audit",
        "period": "Year ended 31 December 2025",
        "preparer": "A. Auditor",
        "workpaper_ref": "C.3.1",
        "prepared_date": "2026-03-31",
        "test_name": "Invoices above approval limit",
        "purpose": "Show that invoices above 1,500 were approved.",
        "procedure": "Filtered invoices >= 1,500.",
        "criteria": "Amount is at least 1,500.00",
        "amount_column": "Amount",
        "app_version": "0.1.0",
    }
    base.update(kw)
    return WorkpaperSpec(**base)
