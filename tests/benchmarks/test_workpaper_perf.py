"""Workpaper memory/time budgets (docs/plan/06-workpaper-export.md 6.5).

Opt-in: ``pytest tests/benchmarks/test_workpaper_perf.py --run-benchmarks -s``.
Rows are generated lazily (no pandas); peak RSS is sampled by a thread.
"""

from __future__ import annotations

import os
import threading
import time

import openpyxl
import polars as pl
import psutil

from tests.support.workpaper import FIXED_NOW, make_spec
from trigger_designer.core.excel import write_frame_xlsx
from trigger_designer.core.workpaper import write_workpaper

ROWS = int(os.environ.get("WORKPAPER_BENCH_ROWS", "500000"))
MB = 1024 * 1024


class _Peak:
    def __enter__(self):
        self.proc = psutil.Process()
        self.base = self.proc.memory_info().rss
        self.peak = self.base
        self._stop = threading.Event()
        self._thread = threading.Thread(target=self._sample)
        self._thread.start()
        self.started = time.perf_counter()
        return self

    def _sample(self):
        while not self._stop.is_set():
            self.peak = max(self.peak, self.proc.memory_info().rss)
            time.sleep(0.05)

    def __exit__(self, *_exc):
        self._stop.set()
        self._thread.join()
        self.elapsed = time.perf_counter() - self.started
        self.delta_mb = (self.peak - self.base) / MB


def _frame(columns: int) -> pl.LazyFrame:
    base = pl.LazyFrame({"Amount": pl.int_range(0, ROWS, eager=True).cast(pl.Float64)})
    return base.with_columns(
        *(
            (pl.col("Amount").cast(pl.Int64) % (i + 7)).cast(pl.String).alias(f"c{i}")
            for i in range(columns)
        )
    )


def test_500k_exceptions_stream_with_bounded_memory(tmp_path):
    spec = make_spec(tmp_path, row_limit=1_048_574)
    with _Peak() as peak:
        result = write_workpaper(spec, _frame(7), now=FIXED_NOW)
    print(f"\n500k x 8: {peak.elapsed:.1f}s, +{peak.delta_mb:.0f} MB")
    assert result.exceptions_written == ROWS and result.csv_path is None
    assert peak.delta_mb < 400 and peak.elapsed < 120
    book = openpyxl.load_workbook(result.path, read_only=True)
    assert book["Exceptions"].max_row == ROWS + 2
    cached = openpyxl.load_workbook(result.path, read_only=True, data_only=True)
    assert cached["Summary"]["B24"].value == ROWS


def test_over_limit_writes_the_full_csv_sidecar(tmp_path):
    spec = make_spec(tmp_path, row_limit=100_000)
    result = write_workpaper(spec, _frame(3), now=FIXED_NOW)
    assert result.truncated and result.exceptions_written == 100_000
    with open(result.csv_path, encoding="utf-8-sig") as handle:
        assert sum(1 for _ in handle) == ROWS + 1
    book = openpyxl.load_workbook(result.path, read_only=True)
    assert book["Exceptions"].max_row == 100_000 + 2


def test_write_frame_xlsx_500k_x_6_bounded_memory(tmp_path):
    with _Peak() as peak:
        written = write_frame_xlsx(_frame(5), str(tmp_path / "big.xlsx"))
    print(f"\nwrite_frame_xlsx 500k x 6: {peak.elapsed:.1f}s, +{peak.delta_mb:.0f} MB")
    assert written == ROWS and peak.delta_mb < 400 and peak.elapsed < 120
