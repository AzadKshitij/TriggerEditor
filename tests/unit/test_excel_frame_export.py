"""``write_frame_xlsx``: one frame -> one sheet (File Output, Data Preview, Profile)."""

from __future__ import annotations

import datetime as dt
import re
import sys
from pathlib import Path

import openpyxl
import polars as pl
import pytest

from trigger_designer.core.excel import ExcelWriteError, engine, write_frame_xlsx

SRC = Path(__file__).resolve().parents[2] / "src" / "trigger_designer"


def _rows(path):
    ws = openpyxl.load_workbook(path).active
    return [[c.value for c in r] for r in ws.iter_rows()]


def test_core_excel_reexports_write_frame_xlsx():
    from trigger_designer.core.excel.frame_export import write_frame_xlsx as direct

    assert direct is write_frame_xlsx


def test_roundtrip_dates_nulls_floats_bools_nan(tmp_path):
    out = tmp_path / "o.xlsx"
    frame = pl.DataFrame(
        {
            "d": [dt.date(2026, 3, 31), None],
            "ts": [dt.datetime(2026, 3, 31, 9, 30, 5), None],
            "f": [1.5, float("nan")],
            "b": [True, None],
            "s": ["x", None],
        }
    )
    assert write_frame_xlsx(frame, str(out)) == 2
    ws = openpyxl.load_workbook(out).active
    assert ws["A2"].value == dt.datetime(2026, 3, 31)
    assert ws["A2"].number_format == "yyyy-mm-dd"
    assert ws["B2"].value == dt.datetime(2026, 3, 31, 9, 30, 5)
    assert ws["B2"].number_format == "yyyy-mm-dd hh:mm:ss"
    assert ws["A3"].value is None and ws["E3"].value is None
    assert ws["C2"].value == 1.5 and ws["C3"].value == "NaN"
    assert ws["D2"].value is True and ws["D3"].value is None


def test_header_is_bold_frozen_and_filtered(tmp_path):
    out = tmp_path / "o.xlsx"
    write_frame_xlsx(pl.LazyFrame({"a": [1, 2], "b": ["x", "y"]}), str(out))
    ws = openpyxl.load_workbook(out).active
    assert ws["A1"].font.b and ws.freeze_panes == "A2"
    assert ws.auto_filter.ref == "A1:B3"
    assert _rows(out) == [["a", "b"], [1, "x"], [2, "y"]]


def test_strings_are_never_formulas(tmp_path):
    out = tmp_path / "o.xlsx"
    write_frame_xlsx(pl.DataFrame({"a": ["=1+1", "+cmd", "-5.5", "ok"]}), str(out))
    ws = openpyxl.load_workbook(out).active
    cells = [ws.cell(r, 1) for r in range(2, 6)]
    assert [c.data_type for c in cells] == ["s", "s", "s", "s"]
    assert [c.quotePrefix for c in cells] == [True, True, False, False]


def test_sheet_name_is_sanitised(tmp_path):
    out = tmp_path / "o.xlsx"
    write_frame_xlsx(
        pl.DataFrame({"a": [1]}), str(out), sheet_name="Profile: [x]/y*" * 3
    )
    name = openpyxl.load_workbook(out).sheetnames[0]
    assert len(name) <= 31 and not re.search(r"[\[\]:*?/\\]", name)


def test_excel_row_limit_is_enforced_without_writing(tmp_path, monkeypatch):
    monkeypatch.setattr(engine, "MAX_ROWS", 4)
    out = tmp_path / "o.xlsx"
    with pytest.raises(ExcelWriteError) as err:
        write_frame_xlsx(pl.DataFrame({"a": [1, 2, 3, 4]}), str(out))
    assert err.value.code == "TD-OUT-013" and not out.exists()


def test_exactly_at_limit_is_written(tmp_path, monkeypatch):
    monkeypatch.setattr(engine, "MAX_ROWS", 4)
    out = tmp_path / "o.xlsx"
    assert write_frame_xlsx(pl.DataFrame({"a": [1, 2, 3]}), str(out)) == 3


def test_column_limit(tmp_path, monkeypatch):
    monkeypatch.setattr(engine, "MAX_COLS", 2)
    frame = pl.DataFrame({"a": [1], "b": [1], "c": [1]})
    with pytest.raises(ExcelWriteError) as err:
        write_frame_xlsx(frame, str(tmp_path / "o.xlsx"))
    assert err.value.code == "TD-OUT-011"


def test_missing_folder_is_td_out_001(tmp_path):
    with pytest.raises(ExcelWriteError) as err:
        write_frame_xlsx(pl.DataFrame({"a": [1]}), str(tmp_path / "nope" / "o.xlsx"))
    assert err.value.code == "TD-OUT-001"


def test_missing_xlsxwriter_is_td_out_003(tmp_path, monkeypatch):
    monkeypatch.setitem(sys.modules, "xlsxwriter", None)
    with pytest.raises(ExcelWriteError) as err:
        write_frame_xlsx(pl.DataFrame({"a": [1]}), str(tmp_path / "o.xlsx"))
    assert err.value.code == "TD-OUT-003"


def test_failed_publish_keeps_old_target_and_cleans_temp(tmp_path, monkeypatch):
    import tempfile

    before = set(Path(tempfile.gettempdir()).glob("td_xl_*"))
    out = tmp_path / "o.xlsx"
    out.write_bytes(b"old")

    def boom(*_args):
        raise ExcelWriteError("x", code="TD-OUT-002")

    monkeypatch.setattr("trigger_designer.core.excel.frame_export.publish", boom)
    with pytest.raises(ExcelWriteError):
        write_frame_xlsx(pl.DataFrame({"a": [1]}), str(out))
    assert out.read_bytes() == b"old"
    assert set(Path(tempfile.gettempdir()).glob("td_xl_*")) == before


def test_locked_target_is_td_out_002(tmp_path, monkeypatch):
    out = tmp_path / "o.xlsx"

    def denied(*_args):
        raise PermissionError("locked")

    monkeypatch.setattr("os.replace", denied)
    with pytest.raises(ExcelWriteError) as err:
        write_frame_xlsx(pl.DataFrame({"a": [1]}), str(out))
    assert err.value.code == "TD-OUT-002" and "open in Excel" in str(err.value)


def test_exotic_dtypes_do_not_crash(tmp_path):
    out = tmp_path / "o.xlsx"
    frame = pl.DataFrame(
        {
            "lst": [[1, 2], None],
            "dur": [dt.timedelta(seconds=90), None],
            "big": [10**16, 5],
            "tz": [dt.datetime(2026, 1, 1, tzinfo=dt.UTC), None],
            "cat": pl.Series(["a", "b"], dtype=pl.Categorical),
        }
    )
    write_frame_xlsx(frame, str(out))
    ws = openpyxl.load_workbook(out).active
    assert ws["A2"].value == "[1, 2]"
    assert ws["B1"].value == "dur (seconds)" and ws["B2"].value == 90
    assert ws["C2"].value == str(10**16) and ws["C3"].value == 5
    assert ws["D2"].value == dt.datetime(2026, 1, 1)
    assert ws["E2"].value == "a"


def test_empty_frame_with_columns_writes_header_only(tmp_path):
    out = tmp_path / "o.xlsx"
    frame = pl.DataFrame({"a": []}, schema={"a": pl.Int64})
    assert write_frame_xlsx(frame, str(out)) == 0
    assert _rows(out) == [["a"]]


def test_multiple_batches_keep_row_order(tmp_path, monkeypatch):
    from trigger_designer.core.excel import spool

    original = spool.Spooled.batches
    monkeypatch.setattr(
        spool.Spooled, "batches", lambda self, size=3: original(self, 3)
    )
    out = tmp_path / "o.xlsx"
    write_frame_xlsx(pl.LazyFrame({"a": list(range(10))}), str(out))
    assert [r[0] for r in _rows(out)[1:]] == list(range(10))


def test_only_engine_imports_xlsxwriter_and_openpyxl_is_read_only():
    offenders = []
    for path in SRC.rglob("*.py"):
        text = path.read_text(encoding="utf-8")
        reads_only = path.name == "fingerprint.py"  # row counts of input workbooks
        uses_openpyxl = re.search(r"^\s*(import|from)\s+openpyxl", text, re.MULTILINE)
        if uses_openpyxl and (
            not reads_only or "write_only" in text or "Workbook(" in text
        ):
            offenders.append(f"openpyxl in {path}")
        if path.name != "engine.py" and re.search(
            r"^\s*(import|from)\s+xlsxwriter|import_module\(.xlsxwriter.\)",
            text,
            re.MULTILINE,
        ):
            offenders.append(f"xlsxwriter in {path}")
    assert offenders == []


def test_polars_write_excel_is_not_used():
    hits = [
        str(p)
        for p in SRC.rglob("*.py")
        if re.search(r"\.write_excel\(", p.read_text(encoding="utf-8"))
    ]
    assert hits == []


def test_no_ws_write_guessing_calls():
    pattern = re.compile(r"(?<![`\w])ws\.write\(")  # code, not a docstring mention
    hits = [
        str(p)
        for folder in ("excel", "workpaper")
        for p in (SRC / "core" / folder).glob("*.py")
        if pattern.search(p.read_text(encoding="utf-8"))
    ]
    assert hits == []


def test_data_preview_export_uses_write_frame_xlsx():
    text = (SRC / "qt" / "models" / "polars_table_viewer.py").read_text(
        encoding="utf-8"
    )
    assert "write_frame_xlsx" in text and "write_excel" not in text


@pytest.mark.slow
def test_streaming_50k_rows_bounded_memory(tmp_path):
    import threading
    import time

    import psutil

    proc, peak, stop = psutil.Process(), [0], threading.Event()
    base = proc.memory_info().rss

    def sample():
        while not stop.is_set():
            peak[0] = max(peak[0], proc.memory_info().rss)
            time.sleep(0.02)

    thread = threading.Thread(target=sample)
    thread.start()
    try:
        frame = pl.LazyFrame({"a": pl.int_range(0, 50_000, eager=True)}).with_columns(
            pl.col("a").cast(pl.String).alias("s"), (pl.col("a") * 1.5).alias("f")
        )
        assert write_frame_xlsx(frame, str(tmp_path / "big.xlsx")) == 50_000
    finally:
        stop.set()
        thread.join()
    assert peak[0] - base < 150 * 1024 * 1024
