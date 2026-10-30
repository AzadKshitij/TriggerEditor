#!/usr/bin/env python3
"""File Output: the core writer and the generated-code entry point.

The streaming `sink_*` calls require a polars with the streaming-engine
"dtype is unknown" fix (see `_join_union_plan`, which panicked on
polars 1.33.1 and sinks cleanly on 1.44.2+).
"""

import os
import sys

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

sys.path.insert(
    0,
    os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "..", "src")),
)

import polars as pl
import pytest
from qtpy.QtWidgets import QApplication

from trigger_designer.core.shaping import file_output as core

APP = QApplication.instance() or QApplication([])


def _make_content(**overrides):
    """Logic-only content (no scene): get_code reads settings and input_vars."""
    from trigger_designer.qt.widgets.nodes.InOut.file_output import (
        FileOutputContent,
    )

    incoming = overrides.pop("incoming_variable", "var_in")
    content = FileOutputContent.__new__(FileOutputContent)
    content.settings = {**core.DEFAULTS, **overrides}
    content.input_vars = [incoming]
    return content


def _join_union_plan():
    """Join + filter + diagonal concat: panicked under sink on polars 1.33.1."""
    orders = pl.LazyFrame({"oid": ["o1", "o2"], "cid": [1, 2], "amt": [10.0, 20.0]})
    cust = pl.LazyFrame({"cid": [1, 3], "name": ["a", "c"]})
    joined = orders.with_columns(pl.lit(1).alias("__l")).join(
        cust.with_columns(pl.lit(1).alias("__r")),
        left_on=["cid"],
        right_on=["cid"],
        how="full",
        coalesce=True,
    )
    matched = joined.filter(
        pl.col("__l").is_not_null() & pl.col("__r").is_not_null()
    ).select(["oid", "cid", "amt", "name"])
    left_only = joined.filter(pl.col("__r").is_null()).select(["oid", "cid", "amt"])
    return pl.concat([matched, left_only], how="diagonal_relaxed")


def test_generated_code_uses_the_core_writer() -> None:
    code = _make_content().get_code()
    assert "trigger_designer.core.shaping.file_output" in code
    assert "write_node" in code
    assert "collect().write_csv" not in code


def test_no_input_fails_loudly() -> None:
    code = _make_content(incoming_variable="").get_code()
    assert "ValueError" in code


def test_live_runner_never_writes(tmp_path) -> None:
    out = tmp_path / "live.csv"
    outs = core.run_node([pl.DataFrame({"a": [1]}).lazy()], {"filePath": str(out)})
    assert outs == []
    assert not out.exists()


def test_empty_path_is_an_error() -> None:
    with pytest.raises(ValueError):
        core.write_node([pl.DataFrame({"a": [1]})], {})


def test_csv_delimiter_and_bom(tmp_path) -> None:
    out = tmp_path / "out.csv"
    core.write_node(
        [pl.DataFrame({"a": [1], "b": [2]})],
        {"filePath": str(out), "delimiter": ";", "include_bom": True},
    )
    raw = out.read_bytes()
    assert raw.startswith(b"\xef\xbb\xbf")
    assert b"a;b" in raw


def test_other_formats_write(tmp_path) -> None:
    frame = pl.DataFrame({"a": [1, 2]})
    cases = {
        "parquet": ("out.parquet", pl.read_parquet),
        "json": ("out.json", pl.read_json),
        "ndjson": ("out.ndjson", pl.read_ndjson),
        "ipc": ("out.ipc", pl.read_ipc),
    }
    for file_format, (name, reader) in cases.items():
        out = tmp_path / name
        core.write_node([frame], {"filePath": str(out), "file_format": file_format})
        assert out.exists(), file_format
        assert reader(out)["a"].to_list() == [1, 2], file_format


def test_custom_delimited_uses_separator(tmp_path) -> None:
    out = tmp_path / "out.txt"
    core.write_node(
        [pl.DataFrame({"a": [1], "b": [2]})],
        {
            "filePath": str(out),
            "file_format": "custom_delimited",
            "delimiter": "|",
        },
    )
    assert "a|b" in out.read_text(encoding="utf-8")


def test_input_is_not_modified(tmp_path) -> None:
    frame = pl.DataFrame({"a": [1, 2]})
    before = frame.clone()
    core.write_node([frame], {"filePath": str(tmp_path / "x.csv")})
    assert frame.equals(before)


def test_generated_csv_code_writes_join_union_plan(tmp_path) -> None:
    """Exec the emitted code against the former panicking plan."""
    out = str(tmp_path / "out.csv")
    code = _make_content(filePath=out).get_code()
    scope = {"var_in": _join_union_plan()}
    exec(compile(code, "<file_output>", "exec"), scope)  # noqa: S102
    result = pl.read_csv(out)
    assert result.shape == (2, 4)
    assert result["oid"].to_list() == ["o1", "o2"]


def test_generated_csv_code_writes_plain_plan(tmp_path) -> None:
    out = str(tmp_path / "plain.csv")
    code = _make_content(filePath=out).get_code()
    scope = {"var_in": pl.LazyFrame({"a": [1, 2]})}
    exec(compile(code, "<file_output>", "exec"), scope)  # noqa: S102
    assert pl.read_csv(out)["a"].to_list() == [1, 2]


def test_windows_path_survives_generated_code(tmp_path) -> None:
    out = str(tmp_path / "win path.csv").replace("/", "\\")
    code = _make_content(filePath=out).get_code()
    scope = {"var_in": pl.LazyFrame({"a": [1]})}
    exec(compile(code, "<file_output>", "exec"), scope)  # noqa: S102
    assert os.path.exists(out)


def test_check_output_header_puts_src_on_path(tmp_path, monkeypatch) -> None:
    from trigger_designer.qt.design_window import CHECK_OUTPUT_HEADER

    (tmp_path / "src").mkdir()
    fake_script = str(tmp_path / "check_output.py")
    monkeypatch.setattr(sys, "path", list(sys.path))
    exec(compile(CHECK_OUTPUT_HEADER, "<header>", "exec"), {"__file__": fake_script})  # noqa: S102
    assert str(tmp_path / "src") in sys.path


def test_excel_format_writes_a_readable_workbook(tmp_path):
    """Excel output needs xlsxwriter; it was undeclared, so this used to fail."""
    import polars as pl
    import pytest

    from trigger_designer.core.shaping import file_output

    pytest.importorskip("xlsxwriter")
    openpyxl = pytest.importorskip("openpyxl")
    path = tmp_path / "out.xlsx"
    frame = pl.DataFrame({"a": [1, 2, None], "b": ["x", None, "z"]}).lazy()
    file_output.write_node([frame], {"filePath": str(path), "file_format": "excel"})
    sheet = openpyxl.load_workbook(path).active
    assert [[c.value for c in row] for row in sheet.iter_rows()] == [
        ["a", "b"],
        [1, "x"],
        [2, None],
        [None, "z"],
    ]


def test_xlsxwriter_is_a_declared_dependency():
    import tomllib
    from pathlib import Path

    project = tomllib.loads(
        (Path(__file__).resolve().parents[2] / "pyproject.toml").read_text("utf-8")
    )
    declared = " ".join(project["project"]["dependencies"])
    assert "xlsxwriter" in declared


def test_excel_format_executes_and_writes_readable_workbook(tmp_path) -> None:
    import openpyxl

    out = tmp_path / "out.xlsx"
    code = _make_content(file_format="excel", filePath=str(out)).get_code()
    exec(  # noqa: S102 - the generated code is the unit under test
        code,
        {"__name__": "__main__"},
        {"var_in": pl.LazyFrame({"a": [1, 2], "b": ["x", None]})},
    )
    rows = [
        [c.value for c in r] for r in openpyxl.load_workbook(out).active.iter_rows()
    ]
    assert rows == [["a", "b"], [1, "x"], [2, None]]


def test_excel_format_never_collects_in_generated_text() -> None:
    code = _make_content(file_format="excel", filePath="C:/t/o.xlsx").get_code()
    assert ".write_excel(" not in code and ".collect()" not in code


# -- parquet compression (D1) -------------------------------------------------
def _parquet_codec(path) -> str:
    pq = pytest.importorskip("pyarrow.parquet")
    return pq.ParquetFile(str(path)).metadata.row_group(0).column(0).compression


@pytest.mark.parametrize(
    ("setting", "codec"),
    [
        (None, "SNAPPY"),
        ("snappy", "SNAPPY"),
        ("zstd", "ZSTD"),
        ("gzip", "GZIP"),
        ("lz4", "LZ4"),
        ("uncompressed", "UNCOMPRESSED"),
    ],
)
def test_parquet_compression_is_honoured(tmp_path, setting, codec) -> None:
    out = tmp_path / "c.parquet"
    core.write_node(
        [pl.DataFrame({"a": list(range(50))})],
        {"filePath": str(out), "file_format": "parquet", "compression": setting},
    )
    assert pl.read_parquet(out)["a"].to_list() == list(range(50))
    assert _parquet_codec(out) == codec


def test_unknown_parquet_compression_is_refused(tmp_path) -> None:
    with pytest.raises(ValueError, match="compression"):
        core.write_node(
            [pl.DataFrame({"a": [1]})],
            {
                "filePath": str(tmp_path / "x.parquet"),
                "file_format": "parquet",
                "compression": "rot13",
            },
        )


def test_compression_is_ignored_for_other_formats(tmp_path) -> None:
    out = tmp_path / "x.csv"
    core.write_node(
        [pl.DataFrame({"a": [1]})],
        {"filePath": str(out), "compression": "zstd"},
    )
    assert pl.read_csv(out)["a"].to_list() == [1]


def test_dead_use_streaming_key_is_gone_but_tolerated(tmp_path) -> None:
    assert "use_streaming" not in core.DEFAULTS
    out = tmp_path / "x.csv"
    core.write_node(
        [pl.DataFrame({"a": [1]})],
        {"filePath": str(out), "use_streaming": False},
    )
    assert out.exists()


def test_generated_parquet_code_compresses_and_passes_policy(tmp_path) -> None:
    from trigger_designer.core.headless.code_policy import validate_generated_code

    out = str(tmp_path / "g.parquet")
    code = _make_content(
        filePath=out, file_format="parquet", compression="zstd"
    ).get_code()
    validate_generated_code(code, io_paths={out})
    exec(compile(code, "<file_output>", "exec"), {"var_in": pl.LazyFrame({"a": [1]})})  # noqa: S102
    assert _parquet_codec(out) == "ZSTD"


def test_generated_code_uses_safe_literals_for_hostile_text(tmp_path) -> None:
    hostile = "x'); import os; os.system('echo hi'); ('.csv"
    code = _make_content(filePath=hostile, delimiter="\n").get_code()
    import ast

    ast.parse(code)
    assert code.count("\n") <= 3  # still a single call, value stayed a value


# -- .xls extension (D2) ------------------------------------------------------
def test_xls_path_is_refused_with_a_plain_td_out_code(tmp_path) -> None:
    from trigger_designer.core.excel.errors import ExcelWriteError

    out = tmp_path / "old.xls"
    with pytest.raises(ExcelWriteError) as caught:
        core.write_node(
            [pl.DataFrame({"a": [1]})],
            {"filePath": str(out), "file_format": "excel"},
        )
    assert caught.value.code == "TD-OUT-012"
    assert ".xlsx" in str(caught.value)
    assert not out.exists()


def test_xls_name_is_refused_for_excel_but_not_for_other_formats(tmp_path) -> None:
    out = tmp_path / "data.xls"
    core.write_node(
        [pl.DataFrame({"a": [1]})],
        {"filePath": str(out), "file_format": "csv"},
    )
    assert out.exists()


def test_generated_code_for_xls_raises_the_plain_error(tmp_path) -> None:
    from trigger_designer.core.excel.errors import ExcelWriteError

    out = tmp_path / "gen.xls"
    code = _make_content(filePath=str(out), file_format="excel").get_code()
    with pytest.raises(ExcelWriteError, match="xlsx"):
        scope = {"var_in": pl.LazyFrame({"a": [1]})}
        exec(compile(code, "<file_output>", "exec"), scope)  # noqa: S102
    assert not out.exists()
